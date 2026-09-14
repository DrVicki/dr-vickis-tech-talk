from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
DOWNLOADS = ROOT / "pages-downloads"
WORKBOOK = DOWNLOADS / "dr-vicki-excel-formula-practice.xlsx"
PROMPT_PACK = DOWNLOADS / "dr-vicki-excel-formula-prompt-pack.zip"

if not WORKBOOK.is_file() or WORKBOOK.stat().st_size == 0:
    raise RuntimeError("Formula practice workbook is missing or empty")
if not PROMPT_PACK.is_file() or PROMPT_PACK.stat().st_size == 0:
    raise RuntimeError("Formula prompt pack is missing or empty")

workbook = load_workbook(WORKBOOK, data_only=False)
expected_sheets = ["Start Here", "Lab 1 - Margin", "Lab 2 - Lookup", "Lab 3 - Status", "Answer Key"]
if workbook.sheetnames != expected_sheets:
    raise RuntimeError(f"Unexpected workbook sheets: {workbook.sheetnames}")
if not workbook["Lab 1 - Margin"]["D6"].value.startswith("=IF("):
    raise RuntimeError("Margin reference formula is missing")
if workbook["Lab 2 - Lookup"]["B9"].value != "Not found":
    raise RuntimeError("Lookup missing-key case is missing")
if workbook["Lab 3 - Status"]["D10"].value != "Due today":
    raise RuntimeError("Status due-today case is missing")
if not workbook["Answer Key"]["B7"].value.startswith("Formula: =IFERROR(XLOOKUP"):
    raise RuntimeError("Lookup answer is missing")
workbook.close()

expected_files = sorted([
    "README.md",
    "01-translate-goal-to-formula.txt",
    "02-stress-test-edge-cases.txt",
    "03-build-reusable-formula-library.txt",
])
with ZipFile(PROMPT_PACK) as archive:
    if archive.testzip() is not None:
        raise RuntimeError("Formula prompt pack is corrupt")
    if sorted(archive.namelist()) != expected_files:
        raise RuntimeError(f"Unexpected prompt pack contents: {archive.namelist()}")

print(f"Formula practice downloads passed: {WORKBOOK.stat().st_size} byte workbook, {PROMPT_PACK.stat().st_size} byte prompt pack")
