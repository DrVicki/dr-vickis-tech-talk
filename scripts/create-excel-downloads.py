from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
DOWNLOADS = ROOT / "pages-downloads"
PROMPTS = DOWNLOADS / "excel-ai-prompts"
WORKBOOK_PATH = DOWNLOADS / "dr-vicki-excel-ai-checklist.xlsx"
ZIP_PATH = DOWNLOADS / "dr-vicki-excel-ai-prompt-pack.zip"

NAVY = "071A2E"
BLUE = "315F95"
CORAL = "FF6048"
LIME = "C7DD2B"
IVORY = "F8F3E8"
WHITE = "FFFFFF"
MUTED = "EDE8DE"
TEXT = "445267"
LIGHT_BLUE = "DCE8F4"

thin = Side(style="thin", color="D9D4CA")
medium_navy = Side(style="medium", color=NAVY)


def style_title(ws, title, subtitle):
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A5"
    ws.merge_cells("A1:F1")
    ws["A1"] = title
    ws["A1"].font = Font(name="Aptos Display", size=24, bold=True, color=WHITE)
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 38

    ws.merge_cells("A2:F2")
    ws["A2"] = subtitle
    ws["A2"].font = Font(name="Aptos", size=11, italic=True, color=TEXT)
    ws["A2"].fill = PatternFill("solid", fgColor=IVORY)
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[2].height = 34

    ws.merge_cells("A3:F3")
    ws["A3"] = "Dr. Vicki's Tech Talk  •  Observe → Verify → Decide → Record"
    ws["A3"].font = Font(name="Aptos", size=10, bold=True, color=NAVY)
    ws["A3"].fill = PatternFill("solid", fgColor=LIME)
    ws["A3"].alignment = Alignment(vertical="center")
    ws.row_dimensions[3].height = 24


def style_table(ws, headers, widths):
    for index, header in enumerate(headers, start=1):
        cell = ws.cell(row=5, column=index, value=header)
        cell.font = Font(name="Aptos", size=10, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = Border(bottom=medium_navy)
    ws.row_dimensions[5].height = 32
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width
    ws.auto_filter.ref = f"A5:{get_column_letter(len(headers))}5"


def style_rows(ws, start_row, end_row, columns):
    for row in range(start_row, end_row + 1):
        fill = PatternFill("solid", fgColor=WHITE if row % 2 == 0 else IVORY)
        for col in range(1, columns + 1):
            cell = ws.cell(row=row, column=col)
            cell.fill = fill
            cell.font = Font(name="Aptos", size=10, color=TEXT)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = Border(bottom=thin)
        ws.row_dimensions[row].height = 56


def add_status_validation(ws, cell_range):
    validation = DataValidation(type="list", formula1='"Not started,In progress,Complete,Needs review"', allow_blank=True)
    validation.error = "Choose a status from the list."
    validation.errorTitle = "Invalid status"
    validation.prompt = "Select the current review status."
    validation.promptTitle = "Workflow status"
    ws.add_data_validation(validation)
    validation.add(cell_range)


def set_print_layout(ws, print_area, landscape=True):
    ws.print_area = print_area
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = 0.25
    ws.page_margins.right = 0.25
    ws.page_margins.top = 0.35
    ws.page_margins.bottom = 0.35
    ws.oddFooter.center.text = "Dr. Vicki's Tech Talk  •  Excel AI Decision Checklist"
    ws.oddFooter.center.size = 8
    ws.oddFooter.center.color = TEXT


def build_workbook():
    wb = Workbook()
    overview = wb.active
    overview.title = "Start Here"
    style_title(
        overview,
        "Excel AI Decision Checklist",
        "Use this workbook with the three downloadable prompt templates. Replace the examples with your own approved data and decision context.",
    )
    style_table(overview, ["Step", "Action", "Suggested question", "Evidence to verify", "Status", "Notes"], [9, 28, 38, 34, 16, 30])
    rows = [
        ("1", "Open a familiar spreadsheet", "What decision or question should this sheet help answer?", "File owner, reporting period, units, filters, and source system", "Not started", ""),
        ("2", "Ask for high-level trends", "What are the three biggest trends in this data?", "Rows and columns supporting each trend; sample size; missing values", "Not started", ""),
        ("3", "Verify the output", "Which source cells support the most important claim?", "Recalculate one claim; inspect outliers; confirm percentages and denominators", "Not started", ""),
        ("4", "Follow the surprise", "What explains the result I did not expect?", "Prior periods, segment mix, seasonality, data-entry changes, and business context", "Not started", ""),
        ("5", "Turn the finding into action", "What should change, what is the trade-off, and what would change your mind?", "Decision owner, constraint, downside, and one critical assumption", "Not started", ""),
        ("6", "Record and repeat", "Which three or four questions should run when this sheet updates?", "Source period, assumptions, approval, next review date, and what changed", "Not started", ""),
    ]
    for row_index, values in enumerate(rows, start=6):
        for col_index, value in enumerate(values, start=1):
            overview.cell(row=row_index, column=col_index, value=value)
    style_rows(overview, 6, 11, 6)
    add_status_validation(overview, "E6:E11")
    set_print_layout(overview, "A1:F11")

    verify = wb.create_sheet("Verification Checklist")
    style_title(
        verify,
        "Verification Checklist",
        "A polished answer or chart can still rest on a weak read. Complete the relevant checks before sharing or acting.",
    )
    style_table(verify, ["Category", "Check", "Why it matters", "Status", "Evidence or cell reference", "Reviewer notes"], [18, 38, 38, 16, 28, 30])
    checks = [
        ("Scope", "Confirm the reporting period and comparison period.", "A valid pattern can be misleading when the date window is wrong."),
        ("Scope", "Confirm units, currency, and whether values are totals, averages, or rates.", "The same number can imply a different decision under a different unit."),
        ("Filters", "Record active filters, hidden rows, and excluded categories.", "The AI may analyze only the visible or selected data."),
        ("Quality", "Check missing, duplicated, or malformed values.", "Gaps and duplicates can create false trends."),
        ("Outliers", "Test whether one unusual value drives the conclusion.", "One event should not automatically become a general pattern."),
        ("Sample", "Confirm the sample is large and representative enough.", "Small or biased samples do not support broad conclusions."),
        ("Time", "Compare prior periods and consider seasonality.", "A seasonal blip can look like a permanent change."),
        ("Math", "Recalculate one important value from the source cells.", "A simple spot check catches many interpretation and formula errors."),
        ("Percentages", "Verify the denominator behind every percentage.", "A correct numerator with the wrong denominator creates a confident error."),
        ("Traceability", "Ask which rows and columns support each key claim.", "A finding should be traceable to visible evidence."),
        ("Decision", "Name the assumption that would change the recommendation.", "This makes uncertainty actionable instead of decorative."),
        ("Governance", "Confirm the AI tool is approved for this data.", "Sensitive or regulated data requires an authorized environment."),
    ]
    for row_index, (category, check, why) in enumerate(checks, start=6):
        values = (category, check, why, "Not started", "", "")
        for col_index, value in enumerate(values, start=1):
            verify.cell(row=row_index, column=col_index, value=value)
    style_rows(verify, 6, 17, 6)
    add_status_validation(verify, "D6:D17")
    set_print_layout(verify, "A1:F17")

    routine = wb.create_sheet("Recurring Review")
    style_title(
        routine,
        "Recurring Review",
        "Use one row per weekly or monthly review. Preserve the questions and assumptions so the next review starts with context.",
    )
    style_table(routine, ["Review date", "Source period", "Standing question", "Key finding", "Confidence and caveat", "Action / owner / next check"], [16, 20, 38, 38, 34, 38])
    for row in range(6, 26):
        routine.cell(row=row, column=1, value="")
        routine.cell(row=row, column=2, value="")
        routine.cell(row=row, column=3, value="")
        routine.cell(row=row, column=4, value="")
        routine.cell(row=row, column=5, value="")
        routine.cell(row=row, column=6, value="")
    style_rows(routine, 6, 25, 6)
    set_print_layout(routine, "A1:F25")

    questions = wb.create_sheet("Standing Questions")
    style_title(
        questions,
        "Standing Questions",
        "Turn Prompt 3 into a stable review routine. Keep each question specific to the sheet and the decision it supports.",
    )
    style_table(questions, ["Priority", "Question to ask every update", "Why this question matters", "Evidence required", "Owner", "Active?"], [12, 42, 38, 34, 20, 14])
    seed_questions = [
        ("1", "What changed most since the previous update?", "Surfaces the largest movement before chart-building begins.", "Current and prior period values", "", "Yes"),
        ("2", "Which segment is driving the change?", "Separates an overall trend from the groups causing it.", "Segment totals and shares", "", "Yes"),
        ("3", "Is the change outside the normal range?", "Distinguishes a meaningful signal from routine variation.", "Historical range and outliers", "", "Yes"),
        ("4", "What action does the verified finding support?", "Connects analysis to a decision and names the trade-off.", "Recommendation, risk, and critical assumption", "", "Yes"),
    ]
    for row_index, values in enumerate(seed_questions, start=6):
        for col_index, value in enumerate(values, start=1):
            questions.cell(row=row_index, column=col_index, value=value)
    for row in range(10, 16):
        for col in range(1, 7):
            questions.cell(row=row, column=col, value="")
    style_rows(questions, 6, 15, 6)
    yes_no = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
    questions.add_data_validation(yes_no)
    yes_no.add("F6:F15")
    set_print_layout(questions, "A1:F15")

    safety = wb.create_sheet("Read Me")
    safety.sheet_view.showGridLines = False
    safety.column_dimensions["A"].width = 4
    safety.column_dimensions["B"].width = 95
    safety.merge_cells("B2:B3")
    safety["B2"] = "How to use this workbook"
    safety["B2"].font = Font(name="Aptos Display", size=24, bold=True, color=WHITE)
    safety["B2"].fill = PatternFill("solid", fgColor=NAVY)
    safety["B2"].alignment = Alignment(vertical="center")
    safety.row_dimensions[2].height = 26
    safety.row_dimensions[3].height = 26
    notes = [
        (5, "1. Start with the familiar", "Choose a spreadsheet you already use and write down the real decision it should support."),
        (8, "2. Ask before you build", "Use the prompt pack to surface patterns before investing time in formulas, charts, or PivotTables."),
        (11, "3. Verify the evidence", "Complete the relevant checks before treating an AI-generated finding as reliable."),
        (14, "4. Keep the human checkpoint", "Record the decision owner, trade-off, critical assumption, and next review date."),
        (17, "Data safety", "Use only an AI tool approved for the data you are handling. Do not upload confidential, regulated, personal, or client data to an unapproved service."),
    ]
    for row, heading, body in notes:
        safety[f"B{row}"] = heading
        safety[f"B{row}"].font = Font(name="Aptos Display", size=16, bold=True, color=CORAL if heading == "Data safety" else NAVY)
        safety[f"B{row + 1}"] = body
        safety[f"B{row + 1}"].font = Font(name="Aptos", size=11, color=TEXT)
        safety[f"B{row + 1}"].alignment = Alignment(wrap_text=True, vertical="top")
        safety.row_dimensions[row + 1].height = 38
    safety["B21"] = "Companion to “Working with Excel Data Without Knowing a Single Formula” • Dr. Vicki's Tech Talk • 2026"
    safety["B21"].font = Font(name="Aptos", size=9, italic=True, color=BLUE)
    set_print_layout(safety, "A1:B21", landscape=False)

    DOWNLOADS.mkdir(parents=True, exist_ok=True)
    wb.save(WORKBOOK_PATH)

    # Reopen to verify workbook integrity and expected sheet structure.
    verified = load_workbook(WORKBOOK_PATH)
    expected = ["Start Here", "Verification Checklist", "Recurring Review", "Standing Questions", "Read Me"]
    if verified.sheetnames != expected:
        raise RuntimeError(f"Unexpected workbook sheets: {verified.sheetnames}")
    verified.close()


def build_zip():
    files = [
        PROMPTS / "README.md",
        PROMPTS / "01-spot-trends-and-assess-risk.txt",
        PROMPTS / "02-turn-findings-into-action.txt",
        PROMPTS / "03-build-a-reusable-routine.txt",
    ]
    with ZipFile(ZIP_PATH, "w", compression=ZIP_DEFLATED) as archive:
        for file in files:
            archive.write(file, arcname=file.name)


if __name__ == "__main__":
    build_workbook()
    build_zip()
    print(f"Created {WORKBOOK_PATH} ({WORKBOOK_PATH.stat().st_size} bytes)")
    print(f"Created {ZIP_PATH} ({ZIP_PATH.stat().st_size} bytes)")
